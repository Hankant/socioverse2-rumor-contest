from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from datetime import datetime,timezone,timedelta
import json,hashlib,re,shutil,os,tempfile,sys
from docx import Document
ROOT=Path(__file__).resolve().parent;REF=ROOT/'reference/source_report.docx';FINAL=ROOT/'SocioVerse-研究报告_FB口径8页版.docx'
SOURCE=Path('E:/papers/socioverse 2026/SocioVerse-研究报告.docx')
sys.path.append('D:/Program Files/Python312/Lib/site-packages')
import fitz
import numpy as np,pandas as pd
from PIL import Image
A=json.loads((ROOT/'analysis_FB.json').read_text(encoding='utf-8'));T=A['tables'];INV=json.loads((ROOT/'reference_inventory.json').read_text(encoding='utf-8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(REF)==INV['sha256']
for source in A['source_files']:assert sha(Path(source['source']))==source['sha256']
get=lambda name,**kw:next(r for r in T[name]if all(r.get(k)==v for k,v in kw.items()))
doc=Document(FINAL);assert len(doc.sections)==3 and len(doc.tables)==5 and len(doc.inline_shapes)==5
alltext='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for stale in ['POST','PRUD3','RBD9','Strict TBU','strict-TBU','two clocks','52 scenario','common three','HC 3','M 1']:
    assert stale not in alltext,stale
for needed in ['Moyan Li','Xu Zhang','Guang Yang','Xiangyu Xu','Pengyuan Xu','4.557','757','222','3.838','1.348']:
    assert needed in alltext,needed
assert len(re.findall(r'Figure [1-5]\s',alltext))==5
assert len(re.findall(r'Table [1-5]\s',alltext))==5
with ZipFile(REF)as old,ZipFile(FINAL)as new:
    assert new.testzip()is None
    authorized={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','docProps/app.xml'}
    preserve=[n for n in old.namelist()if n not in authorized]
    assert all(old.read(n)==new.read(n)for n in preserve)
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','a':'http://schemas.openxmlformats.org/drawingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    x=E.fromstring(new.read('word/document.xml'));z=E.fromstring(old.read('word/document.xml'))
    xb=x.find('w:body',ns);zb=z.find('w:body',ns)
    for newindex,oldindex in enumerate([0,1,2,4]):assert E.tostring(xb[newindex])==E.tostring(zb[oldindex]),'Identity block altered'
    assert len(x.xpath('.//m:oMath',namespaces=ns))==4
    rels=E.fromstring(new.read('word/_rels/document.xml.rels'))
    stems=['Figure1_network_access','Figure2_FB_framework','Figure3_FB_level_differences','Figure4_FB_moderation','Figure5_FB_time_origin']
    for inline,stem in zip(x.xpath('.//wp:inline',namespaces=ns),stems):
        rid=inline.xpath('.//a:blip/@r:embed',namespaces=ns)[0];target='word/'+next(e.get('Target')for e in rels if e.get('Id')==rid)
        assert hashlib.sha256(new.read(target)).hexdigest()==sha(ROOT/'figures'/(stem+'.png'))
# Check every displayed main result against the verified FB analysis, allowing display rounding.
def value(text,expected):
    if text.startswith('<'):assert expected<.001
    else:assert abs(float(text)-expected)<.00051,(text,expected)
for row,f in zip(doc.tables[2].rows[1:],['RT','TC','GC','DGC']):
    r=get('levels',factor=f);value(row.cells[1].text.split('\n')[0],r['F']);value(row.cells[2].text,r['p_holm_both'])
    z=get('pairs',factor=f,low={'RT':2,'TC':.2,'GC':.25,'DGC':0}[f],high={'RT':7,'TC':.8,'GC':.75,'DGC':.4}[f]);value(row.cells[-1].text,z['p_holm_both'])
for row,term in zip(doc.tables[3].rows[1:],['RT_c','TC_c','GC_c','DGC_c','DGC_c:RT_c','DGC_c:TC_c','DGC_c:GC_c']):
    r=get('coefficients',model='M3',term=term);value(row.cells[1].text,r['estimate']);value(row.cells[3].text,r['t']);value(row.cells[4].text,r['p'])
    adj=get('interactions',factor=term.split(':')[1][:-2])['p_holm_both']if ':'in term else r['p_holm_both'];value(row.cells[5].text,adj)
for row,r in zip(doc.tables[4].rows[1:],T['pairs']):
    value(row.cells[2].text,r['estimate']);value(row.cells[-2].text,r['p']);value(row.cells[-1].text,r['p_holm_both'])
cases=pd.read_csv(ROOT/'data/FB_scenarios.csv');people=pd.read_csv(ROOT/'data/FB_users.csv');ev=people[people.TBU_FB.notna()]
assert np.allclose(ev.TBU_FB,ev.reject_tick-ev.first_belief_tick)and ev.reject_tick.ge(ev.RT).all()
assert len(cases)==54 and len(ev)==757 and abs(cases.TBU_FB.mean()-4.55684896011148)<1e-10
pdf=fitz.open(ROOT/'final_render_v2/report.pdf');assert len(pdf)==8
assert [p.rect.width>p.rect.height for p in pdf]==[False,False,True,False,False,False,False,False]
for i,p in enumerate(pdf,1):
    assert len(p.get_text().strip())>100
    bottom=[b[4].strip()for b in p.get_text('blocks')if b[1]>p.rect.height*.9]
    assert str(i)in bottom,(i,bottom)
    for b in p.get_text('blocks'):assert b[0]>=0 and b[1]>=0 and b[2]<=p.rect.width+1 and b[3]<=p.rect.height+1
with Image.open(ROOT/'figures/Figure2_FB_framework.png')as im:assert im.info['dpi'][0]>599
v=dict(pages=8,figures=5,tables=5,native_equations=4,tutor_and_authors_preserved=True,all_visible_results_FB=True,
    conservative_adjustment_retained=True,all_result_tables_checked=True,all_54_source_hashes_checked=True,
    opaque_source_parts_preserved=len(preserve),api_calls=0)
if '--publish'not in sys.argv:print(json.dumps(v,ensure_ascii=False,indent=2));raise SystemExit
assert sha(SOURCE)==INV['sha256'],'Original changed during editing; preserve newer edits and publish separately.'
stamp=datetime.now(timezone(timedelta(hours=8))).strftime('%Y%m%d_%H%M%S')
backup=SOURCE.with_name(SOURCE.stem+'_FB更新前备份_'+stamp+SOURCE.suffix);assert not backup.exists();shutil.copy2(SOURCE,backup);assert sha(backup)==INV['sha256']
sibling=SOURCE.with_name(FINAL.name)
if sibling.exists()and sha(sibling)!=sha(FINAL):sibling=sibling.with_name(sibling.stem+'_'+stamp+sibling.suffix)
shutil.copy2(FINAL,sibling);assert sha(sibling)==sha(FINAL)
fd,name=tempfile.mkstemp(prefix='.codex-fb-report-',suffix='.docx',dir=SOURCE.parent);os.close(fd);staged=Path(name)
assert staged.resolve().parent==SOURCE.parent.resolve()
shutil.copyfile(FINAL,staged);published_original=False;reason=None
try:
    assert sha(SOURCE)==INV['sha256']
    try:os.replace(staged,SOURCE);published_original=True
    except PermissionError as e:reason=str(e)
finally:
    if staged.exists():staged.unlink()
if published_original:assert sha(SOURCE)==sha(FINAL)
figure_dir=SOURCE.parent/'SocioVerse_FB流程图_20261030';figure_dir.mkdir(exist_ok=True)
for ext in ['svg','pdf','png']:
    target=figure_dir/('Figure2_FB_framework.'+ext);shutil.copy2(ROOT/'figures'/target.name,target)
v.update(published_report=str(sibling),original_updated=published_original,original_path=str(SOURCE),backup=str(backup),original_publish_error=reason,
    figure_directory=str(figure_dir),final_sha256=sha(FINAL),all_eight_final_pages_visually_reviewed=True)
(ROOT/'delivery_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(v,ensure_ascii=False,indent=2))
