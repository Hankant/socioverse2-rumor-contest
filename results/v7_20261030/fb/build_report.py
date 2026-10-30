from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import copy,json,hashlib
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as E
ROOT=Path(__file__).resolve().parent;REF=ROOT/'reference/source_report.docx'
INV=json.loads((ROOT/'reference_inventory.json').read_text(encoding='utf-8'));assert hashlib.sha256(REF.read_bytes()).hexdigest()==INV['sha256']
A=json.loads((ROOT/'analysis_FB.json').read_text(encoding='utf-8'));T=A['tables']
get=lambda name,**kw:next(r for r in T[name] if all(r.get(k)==v for k,v in kw.items()))
pv=lambda p:'<.001' if p<.001 else f'{p:.3f}'
ci=lambda r:f"[{r['ci_low']:.3f}, {r['ci_high']:.3f}]"
doc=Document(REF);old=[copy.deepcopy(e)for e in doc._element.body]
for e in list(doc._element.body):
    if e.tag!=qn('w:sectPr'):doc._element.body.remove(e)
basegeo=INV['sections'][0];PAGES=[];CURRENT=[]
def spaced(text):
    if text.startswith('['):return text
    # Preserve identifiers such as HC3, M1--M3 and V7 while fitting ordinary prose.
    text=re.sub(r'([A-Za-z_]+)(?=\d)',lambda m:m.group(1) if m.group(1) in ['HC','M','V','v','SocioVerse'] else m.group(1)+' ',text)
    text=re.sub(r'(?<=[A-Za-z])(?=\.\d)',' ',text)
    text=re.sub(r'(?<=\d)(?=[A-Za-z])',' ',text)
    text=re.sub(r'(?<=%)(?=[A-Za-z])',' ',text)
    text=re.sub(r',(?=[A-Za-z])',', ',text)
    text=re.sub(r'([;:])(?=\S)',r'\1 ',text)
    text=re.sub(r'(?<=[A-Za-z%]),(?=\d)',', ',text)
    text=re.sub(r'(?<=[A-Za-z0-9])(?=\[\d)',' ',text)
    text=re.sub(r'\b(network|targets|ordering|trust|selection)(?=\()',r'\1 ',text)
    return text
def ff(r,size=11,bold=False,italic=False):
    r.font.name='Times New Roman';r.font.size=Pt(size);r.font.bold=bold;r.font.italic=italic;r.font.color.rgb=RGBColor(0,0,0)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SimSun')
def p(text,style=None,size=11):
    text=spaced(text)
    a=doc.add_paragraph(style=style);a.paragraph_format.space_before=Pt(0);a.paragraph_format.space_after=Pt(5);a.paragraph_format.line_spacing=1.0583333333
    if style=='Caption':a.paragraph_format.line_spacing=1;size=9
    ff(a.add_run(text),size);CURRENT.append(text);return a
def h(text,level=1):
    a=doc.add_heading(text,level);a.paragraph_format.space_before=Pt(8);a.paragraph_format.space_after=Pt(5)
    for r in a.runs:ff(r,13 if level==1 else 11.5,True)
    CURRENT.append(text);return a
def page(title=None,orientation=None):
    global CURRENT
    if CURRENT:PAGES.append(CURRENT)
    CURRENT=[]
    if PAGES:
        if orientation:
            s=doc.add_section(WD_SECTION_START.NEW_PAGE)
            for k,v in basegeo.items():setattr(s,k,v)
            if orientation=='landscape':s.page_width=basegeo['page_height'];s.page_height=basegeo['page_width']
            # Continuous numbering and a shared original footer.
            n=s._sectPr.find(qn('w:pgNumType'))
            if n is not None:n.attrib.pop(qn('w:start'),None)
        else:doc.add_page_break()
    if title:h(title)
def figure(stem,caption,width=6.75):
    a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_after=Pt(0);a.paragraph_format.keep_with_next=True
    a.add_run().add_picture(str(ROOT/'figures'/(stem+'.png')),width=Inches(width));c=p(caption,'Caption')
    if width>8:c.paragraph_format.left_indent=Inches(.476);c.paragraph_format.right_indent=Inches(.476)
def table(title,headers,rows,widths,note='',size=9.4):
    a=p(title,'Caption');a.paragraph_format.keep_with_next=True
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    for c,w in zip(tb.columns,widths):c.width=Inches(w)
    pr=tb._tbl.tblPr;bd=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
    pr.append(bd);mar=OxmlElement('w:tblCellMar')
    for edge,v in [('top',65),('bottom',65),('left',75),('right',75)]:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar);rh=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(rh)
    for i,vals in enumerate([headers]+rows):
        row=tb.rows[0]if i==0 else tb.add_row()
        for j,(cell,value,w)in enumerate(zip(row.cells,vals,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            a=cell.paragraphs[0];a.paragraph_format.space_after=Pt(0);a.paragraph_format.space_before=Pt(0);a.paragraph_format.line_spacing=1.02
            a.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or title.startswith(('Table 1','Table 2')) else WD_ALIGN_PARAGRAPH.CENTER
            ff(a.add_run(spaced(str(value))),size,i==0)
            if i==0:
                a.paragraph_format.keep_with_next=True;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5ECF1');cell._tc.get_or_add_tcPr().append(sh)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    CURRENT.append(title)
    if note:p(note,'Caption')
def retain_math(index,rename_gap=False):
    e=copy.deepcopy(old[index])
    if rename_gap:
        for t in e.iter(qn('m:t')):
            if t.text:t.text=t.text.replace('D','DGC')
    doc._element.body.insert(len(doc._element.body)-1,e)
    from docx.text.paragraph import Paragraph
    a=Paragraph(e,doc._body);a.paragraph_format.space_before=Pt(0);a.paragraph_format.space_after=Pt(3);a.paragraph_format.line_spacing=1
def mr(text):
    r=OxmlElement('m:r');rp=OxmlElement('m:rPr');s=OxmlElement('m:sty');s.set(qn('m:val'),'p');rp.append(s);r.append(rp);t=OxmlElement('m:t');t.text=text;r.append(t);return r
def sub(base,script):
    z=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));s=OxmlElement('m:sub');s.append(mr(script));z.extend([e,s]);return z
def equation():
    a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_after=Pt(4)
    mp=OxmlElement('m:oMathPara');om=OxmlElement('m:oMath');mp.append(om)
    for node in [sub('TBU','is'),mr(' = '),sub('r','is'),mr(' − '),sub('b','is'),mr(',   '),sub('r','is'),mr(' ≥ '),sub('RT','s'),mr(';   '),sub('TBU','s'),mr(' = mean('),sub('TBU','is'),mr(').')]:om.append(node)
    a._p.append(mp)
# Page1: retain exact title, authors, tutor and affiliations from the current user file.
page()
for i in [0,1,2,4]:doc._element.body.insert(len(doc._element.body)-1,copy.deepcopy(old[i]))
h('Abstract')
p('Government corrections compete with familiar rumor sources within uneven information networks. We examine response timing (RT), targeting capacity (TC), government credibility (GC) and moderation by degree of group cocooning (DGC) using 54 completed SocioVerse V7 scenarios, each containing 300 heterogeneous LLM users and ten synchronous rounds. Time of belief updating (TBU) starts at a user’s first observed belief and ends at the first observed belief-to-rejection transition in or after the government publication round. The 757 qualifying transitions yield 54 scenario means, averaging 4.557 steps. HC3-robust level tests with retained Holm correction identify RT, TC and GC differences (all adjusted p<.001), while DGC is not significant (p=.272). Relative to RT=2, RT=7 corresponds to 3.838 additional steps; increasing GC from .25 to .75 and TC from .20 to .80 corresponds to reductions of 1.348 and .806 steps. None of the three DGC interactions is supported. TBU includes time awaiting official publication and describes observed converters; 222 unresolved first episodes remain right-censored. Single trajectories per setting, synthetic profiles and the absence of a no-response control limit causal and population-level interpretation.')
p('Keywords  Government rumor correction; time of belief updating; group cocooning; government credibility; LLM agents',size=9.5)
h('1 Research background and objectives')
p('On social media, follower ties and reposts shape which rumor claims and official corrections users encounter. Source trust may influence acceptance, while uneven access can prolong belief within particular groups [1,2]. This study compares TBU across controlled response strategies and examines whether the source-trust gap between cocoon groups changes their associations. It addresses Track B on counterfactual intervention and impact assessment through specified simulation scenarios.')
p('The proposal identifies 125 public rumor events, 21,833 comments, 22,359 reposts, 873 rumor posts in seven categories and 265 questionnaire responses as its wider empirical resource base [5]. The present experiment uses one scripted case and synthetic personas. These resources motivate the research but do not establish behavioral calibration of V7 or representativeness of its simulated users.')
p('The selected outcome follows belief from its initial onset. It therefore evaluates total elapsed duration before an observed update after official publication, including any time already spent believing while awaiting the response. Publication, individual exposure to the response and subsequent belief change are recorded as distinct events.')
# Page2: model design and operational definitions.
page('2 Model setting and experimental design')
p('SOURCE publishes a fixed allegation that safety-renovation funds for old residential neighborhoods were redirected toward commercial-street improvements, delaying safety work. A separate GOVERNMENT account denies the redirection after a claimed records review, explains scheduling differences and promises disclosure. No independently inspectable budget records are attached. Both experimental posts remain identical across settings.')
p('All 300 audience users are LLM agents: 48 follow SOURCE inside the cocoon and 252 do not outside it. Each follows eight same-group and two cross-group peers, producing 3,000 peer ties plus 48 SOURCE ties. Only inside users initially read the original; others encounter it through followee reposts. Direct official delivery does not reveal an unseen rumor root. Encountered roots can be revisited. There is no separate recommendation engine.')
figure('Figure1_network_access','Figure 1  Implemented actors and access. Arrows show message delivery. SOURCE and GOVERNMENT are scripted publishing accounts; audience users make LLM decisions.',width=6.6)
table('Table 1  Factor levels and direct-delivery allocation',['Factor','Symbol','Levels','Operational definition'],[
 ['Response timing','RT','2, 4, 7','Official publication and one-time direct delivery round.'],
 ['Targeting capacity','TC','.20, .80','Inside share of 50 recipients: 10/40 or 40/10 inside/outside.'],
 ['Government credibility','GC','.25, .50, .75','Population mean of assigned prior trust in GOVERNMENT.'],
 ['Degree of group cocooning','DGC','0, .20, .40','Inside-minus-outside mean SOURCE trust; network fixed.']],[1.35,.6,1.1,3.7],
 '3×2×3×3=54 settings; ten rounds each. Targets are sampled without replacement without screening current belief. TC=.20 reaches20.83%/15.87% of inside/outside users; TC=.80 reaches83.33%/3.97%.')
p('Individual SOURCE and government trust are fixed 0–1 priors, distinct from the changing 0–100 belief score. SOURCE mean assignments are:',size=10.5)
retain_math(32,True)
p('Symmetric truncated-normal quantiles (scale .10) and seeded permutations assign trust. Increasing DGC lowers population mean SOURCE trust from .500 to .432 and .364 because group sizes differ; DGC is therefore not a pure structural-network manipulation.',size=10.5)
# Page3: user's three-column composition, updated to FB only.
page('3 Integrated research framework','landscape')
figure('Figure2_FB_framework','Figure 2  Model framework and first-belief TBU. Left: the factorial design and tested moderation paths. Center: four actor roles and the synchronous decision cycle. Right: FB level means, DGC interaction contrasts and RT level differences. Contrasts compare high-minus-low effects between DGC=.4 and 0 (RT7−2; TC.8−.2; GC.75−.25). Whiskers are pointwise HC3 95% intervals; adjusted p values retain the previously tested comparison families. Publishers are scripted; all audience users call the LLM.',width=9.3)
# Page4: complete decision and FB measurement mechanisms.
page('4 Agent decisions and TBU measurement','portrait')
p('Synthetic profiles vary in age (18–74), region, occupation, community-project experience, all 16 MBTI types, verification preference, sharing propensity and repetition sensitivity. For each exposed user, MiniMax-M3 receives all 11 assigned conditions, previous score, actual exposures, current material and retained history. It returns an overall score, three component scores and optional comment/repost choices. Only the overall score classifies belief: >50 is believe; ≤50 is reject. Unjudged is a null baseline for users not yet assessed.')
p('Higher SOURCE trust is prompted to strengthen source persuasiveness. Government trust weights the official response only after actual exposure. Individual values below1/3, from1/3 to below2/3 and at least2/3 receive cautious, conditional and more favorable descriptions of government competence and reliability. These are modeling assumptions; the program does not impose a fixed trust-weighted score or mandatory point change.')
p('For repeat readings beyond the first, n, the prompt uses a saturating familiarity signal. With repetition sensitivity ρ, verification preference v and baseline sharing propensity p:')
retain_math(49);retain_math(50)
p('The coefficients are study settings. History covers the previous three rounds; scores, cumulative exposures and repost flags persist. Each accessible root shows its latest30 earlier comments, including non-followee comments. A user may comment on one visible root and repost one root per round; each root can be reposted only once per user. New actions appear in the next round. Reposts are not independent verification and may express disagreement.')
h('First belief to an observed update',2)
p('For user i in scenario s, b is the first observed score>50; r is the first actual believe-to-reject transition in or after publication round RT. Because publication precedes that round’s decisions, r=RT qualifies. TBU starts at first belief:')
equation()
table('Table 2  Outcome accounting among 16,200 user-scenario instances',['Outcome','Count','Treatment in the TBU analysis'],[
 ['Qualifying observed updates','757','54 scenario means;282 in the publication round and475 later.'],
 ['Unresolved first episodes','222','Right-censored at round10; exclude from the converter mean.'],
 ['Resolved before publication','1','Exclude from the post-publication update endpoint.'],
 ['Never believed','15,220','No belief onset; TBU is not defined.']],[1.65,.55,4.55],
 '980 instances ever believed. Unresolved outcomes are not zero or interpolated. First updating events remain recorded after any later relapse.')
p('Scenario means are equally weighted. Categorical additive models test level differences. Numeric M1 includes RT, TC and GC; M2 adds DGC; M3 adds its interactions with RT, TC and GC. Inputs are centered at RT=13/3, TC=GC=.5 and DGC=.2 without rescaling. HC3 F/t inference uses pointwise95% intervals. FB was retained after exploratory endpoint review, so existing multiplicity families are preserved:8 level tests,20 pairwise contrasts,8 main slopes and6 interactions. Only FB results are presented.',size=10.5)
# Page5: FB factor-level evidence.
page('5 Differences across factor levels')
p('All54 scenarios completed ten rounds (540/540). The mean of scenario TBU values is4.557 steps; pooling the757 observed updates gives4.428 steps. Regression uses the54 scenario means. Adjusted categorical tests support RT, TC and GC differences after the retained Holm correction. DGC has no significant overall level difference.')
figure('Figure3_FB_level_differences','Figure 3  Adjusted first-belief TBU across factor levels. Predictions average equally over the full factorial grid;54 scenario means are observed. Error bars are pointwise HC3 95% intervals. Displayed adjusted p values retain the original eight-test omnibus family.',width=6.75)
rows=[]
for f,lo,hi in [('RT',2,7),('TC',.2,.8),('GC',.25,.75),('DGC',0,.4)]:
    r=get('levels',factor=f);z=get('pairs',factor=f,low=lo,high=hi);ms=sorted([v for v in T['means']if v['factor']==f],key=lambda v:v['level'])
    rows.append([f,f"{r['F']:.3f}\n({int(r['df_num'])},46)",pv(r['p_holm_both']),'\n'.join(f"{v['level']:g}: {v['estimate']:.3f}"for v in ms),f"{z['estimate']:.3f}\n{ci(z)}",pv(z['p_holm_both'])])
table('Table 3  Level tests and high-minus-low contrasts',['Factor','F(df)','Adjusted p','Adjusted means\nlevel:steps','High-minus-low Δ\n95% CI','Pairwise\nadjusted p'],rows,[.5,.85,.8,1.45,2.3,.85],
 'Eight-parameter categorical model; residual df=46. Contrasts:RT7−2,TC.8−.2,GC.75−.25,DGC.4−0. Pairwise p values retain the20-comparison Holm family; intervals are pointwise.',size=9.3)
p('RT=7 rather than2 corresponds to3.838 additional steps (95%CI3.317–4.358). Increasing GC from.25 to.75 corresponds to1.348 fewer steps, and increasing TC from.20 to.80 to.806 fewer steps. These differences describe observed updates within the simulation. The RT contrast includes the period spent awaiting official publication.')
# Page6: moderation and centered numeric response surface.
page('6 Cocooning moderation and effect magnitudes')
fits=[get('fits',model=m)for m in ['M1','M2','M3']];block=get('blocks',block='M3')
p(f"Numeric M1, M2 and M3 explain {fits[0]['R2']*100:.1f}%, {fits[1]['R2']*100:.1f}% and {fits[2]['R2']*100:.1f}% of scenario TBU variation. Adding three DGC interactions increases R² by{block['delta_R2']:.3f}, but the joint HC3 test is not significant:F(3,46)={block['F']:.3f},p={pv(block['p'])}. None of the three interaction terms is significant after the retained correction.")
figure('Figure4_FB_moderation','Figure 4  FB conditional effects and moderation uncertainty. First three panels show RT+1 step, TC.80−.20 and GC+.25 at each DGC level. The forest compares high-minus-low effects between DGC=.4 and0. All intervals are pointwise HC3 95%; interaction p values retain the six-test Holm family.',width=6.75)
rows=[]
for term in ['RT_c','TC_c','GC_c','DGC_c','DGC_c:RT_c','DGC_c:TC_c','DGC_c:GC_c']:
    r=get('coefficients',model='M3',term=term)
    if ':'in term:
        f=term.split(':')[1][:-2];label=f+' × DGC';adj=get('interactions',factor=f)['p_holm_both']
    else:label=term[:-2];adj=r['p_holm_both']
    rows.append([label,f"{r['estimate']:.3f}",ci(r),f"{r['t']:.3f}",pv(r['p']),pv(adj)])
table('Table 4  Centered numeric M3 estimates for FB TBU',['Term','b','95% CI','t(46 df)','Raw p','Adjusted p'],rows,[1.25,.7,2.2,.9,.8,.9],
 'N=54. Main slopes evaluated at DGC=.20 and other centered design means. TC/GC/DGC coefficients are per1.0-unit change. Main slopes retain the eight-test Holm family; interactions the six-test family. Categorical F and numeric slope tests have different null hypotheses.',size=9.4)
z=get('did',factor='GC')
p(f"The GC.75−.25 effect differs by{z['estimate']:.3f} steps between DGC=.4 and0 (95%CI{ci(z)};adjusted p={pv(z['p_holm_both'])}). Its interval includes zero. RT and TC moderation are also uncertain, and fully categorical interaction tests yield no corrected evidence. Separate significance within one DGC stratum does not establish a difference between strata.")
# Page7: complete FB comparisons and substantive interpretation.
page('7 Pairwise comparisons and timing interpretation')
rows=[]
for r in T['pairs']:
    rows.append([r['factor'],f"{r['high']:g} − {r['low']:g}",f"{r['estimate']:.3f}",ci(r),pv(r['p']),pv(r['p_holm_both'])])
table('Table 5  All pairwise level contrasts for FB TBU',['Factor','High−low','Difference\nsteps','95% CI','Raw p','Adjusted p'],rows,[.6,1.15,1.0,2.1,.9,1.0],
 'All ten FB contrasts are reported; adjusted p retains the previously tested20-comparison family. Negative values indicate shorter TBU at the higher factor level. Confidence intervals are pointwise.',size=9.4)
figure('Figure5_FB_time_origin','Figure 5  Illustrative FB measurement. First belief at1, publication at2 and rejection at5 give TBU=4. The interval includes the wait for publication; this example is not a simulated participant record.',width=6.65)
h('Total duration and response timing',2)
p('For users believing before publication, the eligible endpoint cannot precede RT. Delaying publication can therefore lengthen total TBU even if rejection follows quickly afterward. The RT association concerns total duration from initial belief; it does not isolate the processing speed of a correction after receipt. Publication time and each user’s first official exposure are distinct.')
h('Observed updates and unresolved belief',2)
p('Of980 instances that ever believed,757 (77.2%) produced a qualifying observed update. The222 unresolved first episodes remain censored; they are not assigned invented TBU values. One later relapse does not erase a previously observed first update. Low converter-only TBU can coexist with unresolved believers, so both the event count and follow-up limit must accompany the mean.')
p('Government credibility and more inside-directed delivery are supported strategy dimensions within this simulation. However, GC .50 versus .75 does not survive correction (p=.416), so the overall GC test does not establish every adjacent-level difference. TC represents recipient composition rather than accuracy in identifying current believers. DGC moderation remains unsupported.')
# Page8: complete reproducibility, limits and original references.
page('8 Reproducibility limitations and conclusions')
h('Reproducible implementation',2)
p('The implementation uses SocioVerse2[3] and runtime v7-government-trust-persona-20261027-2[4]. MiniMax-M3 is requested with temperature=1.0,top_p=0.95 and maximum output1,536 tokens; no generation seed is sent. Local seeds fix the network(20261018),targets(20261019),ordering(20261020),government trust(20261022),SOURCE trust(20261023) and population selection(20261024). Frozen profiles,prompts,trust assignments,edges,target lists,responses and committed actions support audit and replay. Re-querying can change outputs.')
p('All users observe the pre-update feed before current-round actions are applied. Users with no visible material remain idle. Score ranges,visible targets,evidence IDs and short application summaries for all11conditions are checked. Such summaries document reported input use rather than reveal internal causal reasoning. Unresolved output or service failures preserve successful responses and pause; no synthetic substitute decision is injected. Full rounds are checkpointed, and exported counts,score histories,exposures and individual transition records permit verification.')
h('Scope of inference and further validation',2)
p('Each parameter cell has one stochastic trajectory. The757 user-scenario updates generate54 scenario means, not757 independent experimental replicates. HC3 uncertainty depends on response-surface and error assumptions and does not remove dependence from shared personas or networks. The endpoint was retained after exploratory review; keeping the original multiplicity families prevents a narrower display from artificially lowering adjusted p values.')
p('The scripted case,low simulated belief,synthetic population and absence of a no-response control limit external and causal interpretation. MBTI supplies profile diversity rather than calibrated susceptibility. DGC changes overall SOURCE trust as well as its group gap. Trust direction is embedded in prompt assumptions,so a simulation association does not independently validate that mechanism in human users. Priority checks are independent replications,a no-response arm,alternative cases and networks,longer observation,and sensitivity to trust wording,memory,comments and the threshold.')
h('Conclusion',2)
p('Under FB TBU, later government publication is associated with longer total elapsed belief duration; greater government credibility and more inside-directed targeting are associated with shorter observed duration. The data do not establish cocooning moderation. Conclusions apply to observed converters in the specified simulation,with unresolved episodes and the fixed ten-round horizon reported explicitly.')
h('References')
for i in range(78,83):
    text=''.join(old[i].itertext())
    text=next(b['text']for b in INV['blocks']if b['index']==i)
    a=p(text,size=9);a.paragraph_format.line_spacing=1;a.paragraph_format.space_after=Pt(3)
PAGES.append(CURRENT)
raw=ROOT/'fb_working.docx';FINAL=ROOT/'SocioVerse-研究报告_FB口径8页版.docx';doc.save(raw)
# Preserve all opaque parts; only authoring/body metadata/figure relations can differ.
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','docProps/app.xml'}
with ZipFile(REF)as zr,ZipFile(raw)as zn,ZipFile(FINAL,'w',ZIP_DEFLATED)as zo:
    oldparts=set(zr.namelist())
    for name in zn.namelist():
        data=zr.read(name)if name in oldparts and name not in editable else zn.read(name)
        if name=='docProps/app.xml':
            x=E.fromstring(data);el=x.find('{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Pages')
            if el is not None:el.text='8'
            data=E.tostring(x,xml_declaration=True,encoding='UTF-8',standalone=True)
        zo.writestr(name,data)
    for name in oldparts-set(zn.namelist()):zo.writestr(name,zr.read(name))
with ZipFile(REF)as zr,ZipFile(FINAL)as zn:
    opaque=[n for n in zr.namelist()if n not in editable]
    assert all(zr.read(n)==zn.read(n)for n in opaque)
(ROOT/'planned_pages.json').write_text(json.dumps(PAGES,ensure_ascii=False,indent=2),encoding='utf-8')
print('Built',FINAL,'planned pages',len(PAGES),'figures',len(doc.inline_shapes),'tables',len(doc.tables))
